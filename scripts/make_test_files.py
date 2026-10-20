# 生成 docqa 的测试用文档：一个 docx、一个 pdf、一个 txt
# 用法：python make_test_files.py <输出目录>
import sys
import os
import zlib

out = sys.argv[1]
os.makedirs(out, exist_ok=True)

# ---------- 1. docx ----------
from docx import Document

doc = Document()
doc.add_heading('docqa 测试文档', level=1)
doc.add_paragraph('这是一份用于验证文档解析链路的中文测试文档。')
doc.add_paragraph('第二段：Spring Boot 4.1.1 + Spring Data JPA + MySQL 8.4。')
doc.add_paragraph('第三段：这段文字应该能被解析出来，进入 RAG 的切分流程。')
doc.save(os.path.join(out, 'sample.docx'))

# ---------- 2. txt ----------
with open(os.path.join(out, 'sample.txt'), 'w', encoding='utf-8') as f:
    f.write('这是纯文本测试文件。\nUTF-8 中文应当原样读出。\n第三行。\n')

# ---------- 3. 最小可用的 pdf（手写 PDF 结构，中文字体不嵌入，用 ASCII 内容） ----------
lines = [
    "docqa test pdf",
    "This is a small PDF used to verify the parse pipeline.",
    "Line three.",
]
content = "BT /F1 14 Tf 50 750 Td 18 TL\n"
for line in lines:
    content += "(" + line + ") Tj T*\n"
content += "ET"

objs = []
objs.append("<< /Type /Catalog /Pages 2 0 R >>")
objs.append("<< /Type /Pages /Kids [3 0 R] /Count 1 >>")
objs.append("<< /Type /Page /Parent 2 0 R /MediaBox [0 0 595 842] "
            "/Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>")
objs.append("<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content))
objs.append("<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")

pdf = b"%PDF-1.4\n"
offsets = []
for i, body in enumerate(objs, start=1):
    offsets.append(len(pdf))
    pdf += ("%d 0 obj\n%s\nendobj\n" % (i, body)).encode("latin-1")

xref_pos = len(pdf)
pdf += ("xref\n0 %d\n" % (len(objs) + 1)).encode("latin-1")
pdf += b"0000000000 65535 f \n"
for off in offsets:
    pdf += ("%010d 00000 n \n" % off).encode("latin-1")
pdf += ("trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n"
        % (len(objs) + 1, xref_pos)).encode("latin-1")

with open(os.path.join(out, 'sample.pdf'), 'wb') as f:
    f.write(pdf)

print("written to", out)
for name in sorted(os.listdir(out)):
    print(" ", name, os.path.getsize(os.path.join(out, name)), "bytes")
